"""Excel Lookup Application - Main Flask App.

Serves a web UI for looking up API key details stored in a multi-sheet Excel workbook.
Column A holds the lookup key (API key name); the remaining columns hold the
details shown on each result card.
"""

import functools
import logging
import os
import tempfile
import time

import openpyxl
from flask import Flask, jsonify, request, send_from_directory

# ---------------------------------------------------------------------------
# App setup
# ---------------------------------------------------------------------------

app = Flask(__name__, static_folder="static", template_folder="templates")

# Max upload size: 10 MB
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

DEFAULT_EXCEL_PATH = os.path.join(os.path.dirname(__file__), "data", "api_keys.xlsx")
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}
MAX_FILE_SIZE_MB = 10

# In-memory store: {sheet_name: {headers, rows, index}}
WORKBOOK_STORE: dict = {}

# Simple lookup cache: (sheet, normalized_key) -> [{...}]
_lookup_cache: dict = {}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _normalize(value) -> str:
    """Return a stripped, lower-cased string for case-insensitive matching."""
    if value is None:
        return ""
    return str(value).strip().lower()


def _invalidate_cache():
    """Clear the lookup cache (called when a new workbook is loaded)."""
    _lookup_cache.clear()
    logger.info("Lookup cache invalidated")


def load_workbook(path: str) -> tuple[bool, str]:
    """Load a workbook into WORKBOOK_STORE.

    Returns (success, message). On success the store is populated with every
    sheet; each sheet maps to a dict with 'headers', 'rows', and 'index'
    (keyed on normalised Column A values).
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        logger.error("Failed to load workbook %s: %s", path, exc)
        return False, f"Could not load workbook: {exc}"

    store: dict = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = []
        headers = None
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx == 0:
                headers = [str(h).strip() if h is not None else "" for h in row]
                continue
            if all(h is None or str(h).strip() == "" for h in row):
                continue
            record: dict = {}
            for col_idx, value in enumerate(row):
                name = headers[col_idx] if col_idx < len(headers) else f"col{col_idx + 1}"
                record[name] = value
            rows.append(record)

        index: dict = {}
        for rec in rows:
            key = rec.get(headers[0] if headers else "")
            normalized = _normalize(key)
            index.setdefault(normalized, []).append(rec)

        store[sheet] = {"headers": headers, "rows": rows, "index": index}

    WORKBOOK_STORE.clear()
    WORKBOOK_STORE.update(store)
    _invalidate_cache()
    logger.info("Loaded workbook '%s' with %d sheet(s)", path, len(wb.sheetnames))
    return True, f"Loaded workbook: {path} ({len(wb.sheetnames)} sheets)"


# ---------------------------------------------------------------------------
# Routes – static
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    """Serve the main HTML page."""
    return send_from_directory("templates", "index.html")


# ---------------------------------------------------------------------------
# Routes – health
# ---------------------------------------------------------------------------

@app.route("/health")
def health():
    """Health-check endpoint used by load balancers / Docker HEALTHCHECK."""
    loaded = bool(WORKBOOK_STORE)
    sheets = list(WORKBOOK_STORE.keys()) if loaded else []
    return jsonify({
        "status": "ok",
        "workbook_loaded": loaded,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "timestamp": time.time(),
    })


# ---------------------------------------------------------------------------
# Routes – API
# ---------------------------------------------------------------------------

@app.route("/api/sheets", methods=["GET"])
def list_sheets():
    """Return the list of sheet names in the active workbook."""
    if not WORKBOOK_STORE:
        return jsonify({"error": "No workbook loaded"}), 404
    return jsonify({"sheets": list(WORKBOOK_STORE.keys())})


@app.route("/api/lookup", methods=["GET"])
def lookup():
    """Look up a record by its Column A value.

    Query params:
        sheet:  sheet name to search
        key:    Column A value (case-insensitive)
    """
    sheet = request.args.get("sheet", "").strip()
    key = request.args.get("key", "").strip()

    if not sheet or not key:
        return jsonify({"error": "Missing 'sheet' and/or 'key' query parameters"}), 400

    if not WORKBOOK_STORE:
        return jsonify({"error": "No workbook loaded"}), 404

    if sheet not in WORKBOOK_STORE:
        return jsonify({"error": f"Sheet '{sheet}' not found"}), 404

    cache_key = (sheet, _normalize(key))
    if cache_key in _lookup_cache:
        logger.debug("Cache hit for %s", cache_key)
        return jsonify({"success": True, "sheet": sheet, "data": _lookup_cache[cache_key], "cached": True})

    index = WORKBOOK_STORE[sheet]["index"]
    matches = index.get(_normalize(key))

    if not matches:
        return jsonify({"error": f"No record found for key '{key}' in sheet '{sheet}'"}), 404

    _lookup_cache[cache_key] = matches
    logger.info("Lookup: sheet='%s' key='%s' → %d result(s)", sheet, key, len(matches))
    return jsonify({"success": True, "sheet": sheet, "data": matches, "cached": False})


@app.route("/api/export", methods=["GET"])
def export_csv():
    """Download all rows of a sheet as a CSV file.

    Query params:
        sheet:  sheet name to export
        key:    (optional) filter to rows matching this Column A value
    """
    import csv
    import io

    sheet = request.args.get("sheet", "").strip()
    key = request.args.get("key", "").strip()

    if not sheet:
        return jsonify({"error": "Missing 'sheet' query parameter"}), 400

    if not WORKBOOK_STORE:
        return jsonify({"error": "No workbook loaded"}), 404

    if sheet not in WORKBOOK_STORE:
        return jsonify({"error": f"Sheet '{sheet}' not found"}), 404

    sheet_data = WORKBOOK_STORE[sheet]
    headers = sheet_data["headers"]

    if key:
        rows = sheet_data["index"].get(_normalize(key), [])
    else:
        rows = sheet_data["rows"]

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        # Convert all values to strings for CSV safety
        writer.writerow({h: ("" if row.get(h) is None else str(row[h])) for h in headers})

    csv_bytes = output.getvalue().encode("utf-8-sig")  # BOM for Excel compatibility

    from flask import Response
    filename = f"{sheet}{'_' + key if key else ''}.csv"
    return Response(
        csv_bytes,
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.route("/api/upload", methods=["POST"])
def upload():
    """Accept a new Excel file, replacing the currently loaded workbook."""
    if "file" not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "No file selected"}), 400

    ext = os.path.splitext(file.filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        return jsonify({"error": f"Only {', '.join(ALLOWED_EXTENSIONS)} files are supported"}), 400

    # Check content-length if provided by client
    content_length = request.content_length
    if content_length and content_length > app.config["MAX_CONTENT_LENGTH"]:
        return jsonify({"error": f"File exceeds the {MAX_FILE_SIZE_MB} MB size limit"}), 413

    tmp_fd = None
    tmp_path = None
    try:
        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
        file.save(tmp_path)

        # Extra server-side size validation
        actual_size = os.path.getsize(tmp_path)
        if actual_size > app.config["MAX_CONTENT_LENGTH"]:
            return jsonify({"error": f"File exceeds the {MAX_FILE_SIZE_MB} MB size limit"}), 413

        success, message = load_workbook(tmp_path)
        if success:
            logger.info("Upload successful: %s (%d bytes)", file.filename, actual_size)
            return jsonify({"success": True, "message": message, "sheets": list(WORKBOOK_STORE.keys())})
        return jsonify({"error": message}), 500

    except Exception as exc:
        logger.exception("Upload failed")
        return jsonify({"error": f"Upload failed: {exc}"}), 500

    finally:
        if tmp_fd is not None:
            try:
                os.close(tmp_fd)
            except Exception:
                pass
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


@app.route("/api/default", methods=["GET"])
def default_workbook():
    """Load the default workbook shipped with the repo (if present)."""
    if not os.path.exists(DEFAULT_EXCEL_PATH):
        return jsonify({"error": f"Default workbook not found: {DEFAULT_EXCEL_PATH}"}), 404

    success, message = load_workbook(DEFAULT_EXCEL_PATH)
    if not success:
        return jsonify({"error": message}), 500
    return jsonify({"success": True, "message": message, "sheets": list(WORKBOOK_STORE.keys())})


# ---------------------------------------------------------------------------
# Error handlers
# ---------------------------------------------------------------------------

@app.errorhandler(413)
def request_entity_too_large(e):
    return jsonify({"error": f"File exceeds the {MAX_FILE_SIZE_MB} MB size limit"}), 413


@app.errorhandler(404)
def not_found(e):
    return jsonify({"error": "Not found"}), 404


@app.errorhandler(500)
def internal_error(e):
    logger.exception("Internal server error")
    return jsonify({"error": "Internal server error"}), 500


# ---------------------------------------------------------------------------
# Startup – auto-load default workbook
# ---------------------------------------------------------------------------

if os.path.exists(DEFAULT_EXCEL_PATH):
    load_workbook(DEFAULT_EXCEL_PATH)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
